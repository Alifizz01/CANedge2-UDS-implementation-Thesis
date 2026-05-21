"""
Generate the Weekly Progress & Thesis Plan Excel file.
Run: python generate_excel.py
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

wb = openpyxl.Workbook()

# ─── Color scheme ───
HEADER_FILL = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True)
DONE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
CURRENT_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
UPCOMING_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
NORMAL_FONT = Font(name="Calibri", size=11)
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="2E5090")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

def style_cell(ws, row, col, bold=False, fill=None, wrap=True):
    cell = ws.cell(row=row, column=col)
    cell.font = BOLD_FONT if bold else NORMAL_FONT
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    cell.border = thin_border
    if fill:
        cell.fill = fill
    return cell

# ═══════════════════════════════════════════════════════════
# SHEET 1: Weekly Progress
# ═══════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Weekly Progress"

# Title
ws1.merge_cells("A1:F1")
title_cell = ws1["A1"]
title_cell.value = "Bachelor Thesis - Weekly Progress Tracker"
title_cell.font = TITLE_FONT
title_cell.alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[1].height = 30

ws1.merge_cells("A2:F2")
ws1["A2"].value = "Inside the Battery: Entschluesselung von Batteriedaten ueber den CAN-Bus"
ws1["A2"].font = Font(name="Calibri", size=11, italic=True, color="666666")

# Headers
headers = ["Week", "Date Range", "Phase", "What I Did", "Deliverables / Output", "Status"]
for c, h in enumerate(headers, 1):
    ws1.cell(row=4, column=c, value=h)
style_header_row(ws1, 4, 6)

# Column widths
ws1.column_dimensions["A"].width = 10
ws1.column_dimensions["B"].width = 22
ws1.column_dimensions["C"].width = 18
ws1.column_dimensions["D"].width = 55
ws1.column_dimensions["E"].width = 40
ws1.column_dimensions["F"].width = 12

# Weekly data
weeks = [
    {
        "week": "Week 1",
        "dates": "03.03 - 08.03.2026",
        "phase": "Setup & First Analysis",
        "tasks": (
            "- Got the MF4 files from the CANedge2 loaded into Python\n"
            "- Wrote mf4_reader.py to parse all 4 log files into a single DataFrame\n"
            "- Built reverse_engineer.py with 9 different analysis techniques\n"
            "  (byte time series, bit heatmaps, boundary detection, 16-bit testing,\n"
            "   counter/checksum detection, ASCII detection, scaling analysis, per-source comparison)\n"
            "- Ran the full analysis pipeline on all 7 CAN IDs (0x065 - 0x06F)\n"
            "- Wrote can_decoder.py with signal definitions for confirmed signals\n"
            "- Tried UDS PID matching - found 145 battery PIDs available but\n"
            "  they use 29-bit CAN IDs (UDS diagnostic) vs our 11-bit passive data\n"
            "- Created full documentation of all findings\n"
            "- Generated 39 RE plots + 8 decoded plots\n"
            "- Exported merged CAN data as CSV\n"
            "- Reorganised folder structure (src/, data/, output/, docs/)"
        ),
        "deliverables": (
            "mf4_reader.py, reverse_engineer.py,\n"
            "can_decoder.py, uds_decoder.py,\n"
            "analyze_patterns.py\n"
            "39 plots in output/plots/\n"
            "DOCUMENTATION.md (full analysis)\n"
            "output/can_data.csv"
        ),
        "status": "Done"
    },
    {
        "week": "Week 2",
        "dates": "09.03 - 15.03.2026",
        "phase": "Tools & Organisation",
        "tasks": (
            "- Built a full CAN Analyzer GUI in Python (PyQt5 + pyqtgraph)\n"
            "  6 tabs: Sniffer, ID Analysis, Signal Plotter, Payload Analysis,\n"
            "  Flow Analysis, Hex Matrix\n"
            "- Python GUI too slow -> built native CAN Analyzer v2 in Rust (egui)\n"
            "  4 tabs: Data Flow, Sniffer, Signal Plotter, Stats\n"
            "- Data Flow tab: real-time hex matrix with byte-change highlighting\n"
            "- Added direct MF4 loading to the Rust app\n"
            "- Created organisation documents (weekly tracker, thesis plan,\n"
            "  project overview)\n"
            "- Refocused thesis on BATTERY parameters per lecturer feedback\n"
            "- Created data collection strategy document\n"
            "- Created data collection procedure template"
        ),
        "deliverables": (
            "src/can_analyzer.py (Python GUI)\n"
            "can_analyzer/ (Rust app, can-analyzer.exe)\n"
            "Weekly_Progress.xlsx\n"
            "Thesis_Plan_and_Methodology.md\n"
            "Data_Collection_Strategy.md\n"
            "Data_Collection_Procedure.md"
        ),
        "status": "Done"
    },
    {
        "week": "Week 3",
        "dates": "16.03 - 22.03.2026",
        "phase": "Data Collection",
        "tasks": (
            "- Check CANedge2 capabilities: can it send 29-bit extended CAN frames?\n"
            "- Test UDS request via CANedge2 transmit list (Strategy A):\n"
            "  send SOC request (03 22 02 8C) to BMS at 0x17FC007B\n"
            "- Check if BMS responds (look for 0x17FE007B in logs)\n"
            "- If no response: check diagnostic session setup (0x10 0x03)\n"
            "- If CANedge2 can't transmit: evaluate Strategy B\n"
            "  (external OBD adapter + CANedge2 as passive logger)\n"
            "- First battery data recording (standstill baseline)\n"
            "- Re-examine existing data with battery lens:\n"
            "  are 0x066/0x067 temperatures actually battery temps?\n"
            "- Try recording during driving for battery current/voltage changes"
        ),
        "deliverables": (
            "Strategy validation: which approach works\n"
            "First battery data (SOC, voltage, current)\n"
            "Updated CANedge2 config\n"
            "Session logs with annotations"
        ),
        "status": "In Progress"
    },
    {
        "week": "Week 4",
        "dates": "23.03 - 29.03.2026",
        "phase": "Battery Data Collection",
        "tasks": (
            "- Full battery data collection using validated strategy\n"
            "- Scenario 1: Standstill baseline (5 min, all battery params stable)\n"
            "- Scenario 2: Constant speed driving (50, 100 km/h)\n"
            "  -> observe battery current/voltage under steady load\n"
            "- Scenario 3: Acceleration & regen braking\n"
            "  -> peak current, regen current direction change\n"
            "- Scenario 4: Extended drive (20-30 min) for SOC decrease tracking\n"
            "- If possible: configure cell voltage requests (ISO-TP needed)\n"
            "- Manual annotation for every test with timestamps"
        ),
        "deliverables": (
            "Battery recordings (4+ scenarios)\n"
            "Annotation logs\n"
            "Initial UDS response decoding"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 5",
        "dates": "30.03 - 05.04.2026",
        "phase": "Battery Data Collection 2",
        "tasks": (
            "- Charging session recording (DC fast charge if available)\n"
            "  -> SOC increase, charging current, battery temp rise\n"
            "- Cold start recording (after overnight park)\n"
            "  -> low battery temp, PTC heater activity, power limits\n"
            "- Temperature under load (mixed driving, watch temp sensors)\n"
            "- If cell voltages accessible: record cell voltage spread\n"
            "- Start writing: CAN bus + UDS theory chapter\n"
            "- Start writing: battery fundamentals section"
        ),
        "deliverables": (
            "Charging + cold start recordings\n"
            "Draft: Theoretical background (~5 pages)"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 6",
        "dates": "06.04 - 12.04.2026",
        "phase": "Battery Analysis",
        "tasks": (
            "- Decode all captured UDS battery responses\n"
            "- SOC: validate against dashboard reading\n"
            "- Voltage: check if pack voltage matches expected range (~350-410V)\n"
            "- Current: verify direction (negative=discharge, positive=charge)\n"
            "- Temperature: compare with ambient temp, check heating behaviour\n"
            "- Cross-correlate: battery current vs driving speed\n"
            "- Map out which PIDs responded and which didn't\n"
            "- Write CANedge2 implementation chapter"
        ),
        "deliverables": (
            "Decoded battery parameter table\n"
            "Validation against dashboard readings\n"
            "Draft: CANedge2 + Implementation (~5 pages)"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 7",
        "dates": "13.04 - 19.04.2026",
        "phase": "Deep Analysis",
        "tasks": (
            "- Cell voltage analysis (if data available):\n"
            "  voltage spread, weakest cell, balancing behaviour\n"
            "- Temperature distribution across 18 sensor points\n"
            "- Battery behaviour comparison: driving vs charging vs idle\n"
            "- Power calculation: P = V * I from decoded voltage and current\n"
            "- Energy consumption analysis: kWh over distance\n"
            "- Build partial DBC file with confirmed battery signals\n"
            "- Re-examine passive CAN data (0x065-0x06F) for any battery signals"
        ),
        "deliverables": (
            "Cell voltage analysis plots\n"
            "Temperature distribution maps\n"
            "vw_id_buzz_battery.dbc (partial)"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 8",
        "dates": "20.04 - 26.04.2026",
        "phase": "Writing",
        "tasks": (
            "- Write methodology chapter:\n"
            "  data collection strategy, why UDS was needed,\n"
            "  9 RE techniques for passive data, controlled test scenarios\n"
            "- Write analysis chapter (first half):\n"
            "  passive CAN analysis of 0x065-0x06F,\n"
            "  UDS battery parameter decoding approach"
        ),
        "deliverables": (
            "Draft: Methodology (~8 pages)\n"
            "Draft: Analysis part 1 (~5 pages)"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 9",
        "dates": "27.04 - 03.05.2026",
        "phase": "Writing",
        "tasks": (
            "- Write analysis chapter (second half):\n"
            "  battery parameter results, SOC/voltage/current/temp\n"
            "  cell voltage analysis, charging vs driving comparison\n"
            "- Create publication-quality figures (300 DPI)\n"
            "  consistent style across all battery data plots\n"
            "- Write results chapter: what was decoded and how confident"
        ),
        "deliverables": (
            "Draft: Analysis part 2 + Results (~10 pages)\n"
            "All final figures"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 10",
        "dates": "04.05 - 10.05.2026",
        "phase": "Validation",
        "tasks": (
            "- Validation data collection: new recordings specifically to verify\n"
            "- Compare decoded SOC vs dashboard SOC (% error)\n"
            "- Compare decoded temperature vs external thermometer\n"
            "- Check voltage plausibility (expected ~3.6-4.2V per cell, ~350-410V pack)\n"
            "- Compute error metrics: MAE, R-squared, max deviation\n"
            "- Document validation results with evidence"
        ),
        "deliverables": (
            "Validation dataset\n"
            "Error metrics report\n"
            "Validation results section"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 11",
        "dates": "11.05 - 17.05.2026",
        "phase": "Writing",
        "tasks": (
            "- Write discussion chapter:\n"
            "  limitations (gateway filtering, CAN bus access, ISO-TP),\n"
            "  comparison with other approaches (PASTA, comma.ai opendbc),\n"
            "  implications for EV battery research,\n"
            "  ethical/security considerations\n"
            "- Write conclusion and future work\n"
            "- Write abstract (English + German)"
        ),
        "deliverables": (
            "Draft: Discussion + Conclusion (~8 pages)\n"
            "Abstract"
        ),
        "status": "Planned"
    },
    {
        "week": "Week 12",
        "dates": "18.05 - 24.05.2026",
        "phase": "Finalisation",
        "tasks": (
            "- Full read-through and proofreading\n"
            "- Check all references and bibliography\n"
            "- Appendix: code listings, DBC file, raw data samples,\n"
            "  UDS PID reference table, all analysis plots\n"
            "- PDF formatting, page numbers, margins\n"
            "- Final review, print, submit"
        ),
        "deliverables": (
            "Final thesis document (45-60 pages)\n"
            "Complete code repository\n"
            "Partial DBC file (battery signals)"
        ),
        "status": "Planned"
    },
]

for i, w in enumerate(weeks):
    row = 5 + i
    ws1.row_dimensions[row].height = max(80, len(w["tasks"].split("\n")) * 15)

    status = w["status"]
    if status == "Done":
        fill = DONE_FILL
    elif status == "In Progress":
        fill = CURRENT_FILL
    else:
        fill = UPCOMING_FILL

    style_cell(ws1, row, 1, bold=True, fill=fill).value = w["week"]
    style_cell(ws1, row, 2, fill=fill).value = w["dates"]
    style_cell(ws1, row, 3, bold=True, fill=fill).value = w["phase"]
    style_cell(ws1, row, 4, fill=fill).value = w["tasks"]
    style_cell(ws1, row, 5, fill=fill).value = w["deliverables"]
    style_cell(ws1, row, 6, bold=True, fill=fill).value = status

# Legend row
legend_row = 5 + len(weeks) + 1
ws1.cell(row=legend_row, column=1, value="Legend:").font = BOLD_FONT
ws1.cell(row=legend_row, column=2, value="Done").fill = DONE_FILL
ws1.cell(row=legend_row, column=2).font = NORMAL_FONT
ws1.cell(row=legend_row, column=3, value="In Progress").fill = CURRENT_FILL
ws1.cell(row=legend_row, column=3).font = NORMAL_FONT
ws1.cell(row=legend_row, column=4, value="Planned").fill = UPCOMING_FILL
ws1.cell(row=legend_row, column=4).font = NORMAL_FONT


# ═══════════════════════════════════════════════════════════
# SHEET 2: Signal Tracking
# ═══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Signal Tracking")

ws2.merge_cells("A1:H1")
ws2["A1"].value = "CAN Signal Tracking"
ws2["A1"].font = TITLE_FONT
ws2.row_dimensions[1].height = 30

sig_headers = ["CAN ID", "Byte(s)", "Signal Name", "Status", "Confidence", "Scaling", "Unit", "Validated?"]
for c, h in enumerate(sig_headers, 1):
    ws2.cell(row=3, column=c, value=h)
style_header_row(ws2, 3, 8)

ws2.column_dimensions["A"].width = 10
ws2.column_dimensions["B"].width = 10
ws2.column_dimensions["C"].width = 20
ws2.column_dimensions["D"].width = 12
ws2.column_dimensions["E"].width = 12
ws2.column_dimensions["F"].width = 14
ws2.column_dimensions["G"].width = 10
ws2.column_dimensions["H"].width = 30

signals = [
    ("0x065", "B[0]", "NM_STATE", "Decoded", "HIGH", "enum", "-", "Need more wake cycles"),
    ("0x066", "B[0]", "OperatingMode", "Decoded", "HIGH", "enum (0x01/0x03)", "-", "Partial"),
    ("0x066", "B[1:2]", "MsgCounter", "Decoded", "HIGH", "1:1 (big-endian)", "count", "Yes, wrapping OK"),
    ("0x066", "B[3]", "Temperature_A", "Decoded", "MEDIUM", "raw - 40", "deg C", "Need thermometer"),
    ("0x066", "B[4:5]", "ASCII Marker 'T-'", "Decoded", "HIGH", "ASCII", "-", "Yes"),
    ("0x067", "B[0]", "Dynamic_A", "Partial", "LOW", "unknown", "?", "Need controlled scenarios"),
    ("0x067", "B[1]", "Temperature_B", "Partial", "MEDIUM", "unknown", "deg C?", "Need validation"),
    ("0x067", "B[3]", "Multiplexer", "Decoded", "MEDIUM", "3-bit enum", "-", "Need more states"),
    ("0x067", "B[5]", "Temperature_C", "Decoded", "MEDIUM", "raw - 40", "deg C", "Cross-validated w/ 0x066"),
    ("0x068", "B[0]", "Signal_D", "Partial", "LOW", "unknown", "?", "Need scenarios"),
    ("0x068", "B[3]", "StatusFlag", "Decoded", "HIGH", "binary 0/1", "-", "Yes"),
    ("0x06A", "ALL", "EventMsg", "Unknown", "LOW", "-", "-", "Only 126 msgs, need more data"),
    ("0x06B", "B[0:3]", "Signals F-H", "Partial", "LOW", "unknown", "?", "Need scenarios"),
    ("0x06F", "B[0:7]", "8 independent signals", "Partial", "LOW", "unknown", "?", "NOT encrypted (proved)"),
]

decoded_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
partial_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
unknown_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

for i, sig in enumerate(signals):
    row = 4 + i
    status = sig[3]
    if status == "Decoded":
        fill = decoded_fill
    elif status == "Partial":
        fill = partial_fill
    else:
        fill = unknown_fill

    for c, val in enumerate(sig, 1):
        cell = style_cell(ws2, row, c, fill=fill)
        cell.value = val

# Legend
lr = 4 + len(signals) + 1
ws2.cell(row=lr, column=1, value="Legend:").font = BOLD_FONT
ws2.cell(row=lr, column=2, value="Decoded").fill = decoded_fill
ws2.cell(row=lr, column=3, value="Partial").fill = partial_fill
ws2.cell(row=lr, column=4, value="Unknown").fill = unknown_fill


# ═══════════════════════════════════════════════════════════
# SHEET 3: Tools & Software
# ═══════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Tools and Software")

ws3.merge_cells("A1:D1")
ws3["A1"].value = "Tools & Software Used"
ws3["A1"].font = TITLE_FONT
ws3.row_dimensions[1].height = 30

tool_headers = ["Tool / Software", "Version", "Purpose", "Status"]
for c, h in enumerate(tool_headers, 1):
    ws3.cell(row=3, column=c, value=h)
style_header_row(ws3, 3, 4)

ws3.column_dimensions["A"].width = 25
ws3.column_dimensions["B"].width = 18
ws3.column_dimensions["C"].width = 50
ws3.column_dimensions["D"].width = 15

tools = [
    ("Python", "3.13", "Main scripting language for analysis", "Installed"),
    ("asammdf", "7.x", "Reading MF4 (ASAM MDF4) log files", "Installed"),
    ("pandas", "2.x", "Data manipulation and DataFrame handling", "Installed"),
    ("numpy", "1.24+", "Numerical operations, array processing", "Installed"),
    ("matplotlib", "3.7+", "Plotting and figure generation", "Installed"),
    ("PyQt5", "5.15+", "Python GUI framework (CAN Analyzer v1)", "Installed"),
    ("pyqtgraph", "0.13+", "Fast interactive plots in Python GUI", "Installed"),
    ("Rust", "1.94", "Native CAN Analyzer v2 (fast, compiled)", "Installed"),
    ("eframe / egui", "0.31", "Rust GUI framework (immediate mode)", "Installed"),
    ("egui_plot", "0.31", "Interactive plotting in Rust app", "Installed"),
    ("CANedge2", "Hardware", "CAN bus data logger (SD card, WiFi)", "Available"),
    ("asammdf GUI", "8.2.5", "Visual MF4 file inspection (portable)", "Available"),
    ("GPS Logger App", "TBD", "Speed ground truth for validation", "Need to install"),
    ("External Thermometer", "Hardware", "Temperature signal validation", "Need to get"),
    ("LaTeX / Word", "TBD", "Thesis document writing", "Need to set up"),
]

for i, tool in enumerate(tools):
    row = 4 + i
    for c, val in enumerate(tool, 1):
        cell = style_cell(ws3, row, c)
        cell.value = val
        if val == "Installed" or val == "Available":
            cell.fill = decoded_fill
        elif val.startswith("Need"):
            cell.fill = partial_fill


# ═══════════════════════════════════════════════════════════
# SHEET 4: Data Overview
# ═══════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Data Overview")

ws4.merge_cells("A1:E1")
ws4["A1"].value = "CAN Data Overview"
ws4["A1"].font = TITLE_FONT
ws4.row_dimensions[1].height = 30

ws4["A3"].value = "Current Data (Initial Recording)"
ws4["A3"].font = SUBHEADER_FONT
ws4["A3"].fill = SUBHEADER_FILL
ws4.merge_cells("A3:E3")

data_headers = ["Property", "Value"]
for c, h in enumerate(data_headers, 1):
    ws4.cell(row=4, column=c, value=h)
style_header_row(ws4, 4, 2)

ws4.column_dimensions["A"].width = 30
ws4.column_dimensions["B"].width = 45

data_info = [
    ("Vehicle", "Volkswagen ID Buzz (MEB platform)"),
    ("Logger", "CANedge2 data logger"),
    ("Data Format", "MF4 (ASAM MDF v4.11)"),
    ("Log Files", "4 sessions (00000002 - 00000005)"),
    ("Total Frames", "9,146 CAN frames"),
    ("Total Duration", "~333 seconds (~5.5 minutes)"),
    ("Unique CAN IDs", "7 (0x065, 0x066, 0x067, 0x068, 0x06A, 0x06B, 0x06F)"),
    ("Bus Channel", "Channel 9"),
    ("Direction", "0 (received)"),
    ("Bit Rate", "500 kbps (standard VW)"),
]

for i, (prop, val) in enumerate(data_info):
    row = 5 + i
    style_cell(ws4, row, 1, bold=True).value = prop
    style_cell(ws4, row, 2).value = val

# CAN ID table
can_row = 5 + len(data_info) + 1
ws4.cell(row=can_row, column=1, value="Per-CAN-ID Breakdown").font = SUBHEADER_FONT
ws4.cell(row=can_row, column=1).fill = SUBHEADER_FILL
ws4.merge_cells(f"A{can_row}:E{can_row}")

can_headers = ["CAN ID", "Messages", "DLC", "Frequency", "Present In"]
for c, h in enumerate(can_headers, 1):
    ws4.cell(row=can_row+1, column=c, value=h)
style_header_row(ws4, can_row+1, 5)

ws4.column_dimensions["C"].width = 8
ws4.column_dimensions["D"].width = 12
ws4.column_dimensions["E"].width = 20

can_ids = [
    ("0x06F", "3,034", "8", "~9 Hz", "All sources"),
    ("0x065", "3,002", "1", "~9 Hz", "All sources"),
    ("0x066", "920", "6", "~5 Hz", "00000004 only"),
    ("0x067", "688", "8", "~4.5 Hz", "00000004 only"),
    ("0x068", "688", "4", "~4.5 Hz", "00000004 only"),
    ("0x06B", "688", "5", "~4.5 Hz", "00000004 only"),
    ("0x06A", "126", "varies", "sporadic", "00000004 only"),
]

for i, can in enumerate(can_ids):
    row = can_row + 2 + i
    for c, val in enumerate(can, 1):
        style_cell(ws4, row, c).value = val


# ═══════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════
out_path = os.path.join(os.path.dirname(__file__), "Weekly_Progress.xlsx")
wb.save(out_path)
print(f"Created: {out_path}")
