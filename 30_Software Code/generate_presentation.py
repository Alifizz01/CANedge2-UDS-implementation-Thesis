"""
Generate PowerPoint presentation about the battery CAN bus reverse engineering research.
Covers: research focus, methodology, strategy, findings, next steps.

Run: cd src && python -X utf8 generate_presentation.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

OUTPUT_DIR = Path(__file__).parent.parent / "10_Literature"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Color scheme
DARK_BLUE = RGBColor(0, 51, 102)
MED_BLUE = RGBColor(0, 100, 180)
LIGHT_BLUE = RGBColor(200, 220, 240)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(30, 30, 30)
GRAY = RGBColor(100, 100, 100)
DARK_GRAY = RGBColor(60, 60, 60)
GREEN = RGBColor(0, 130, 60)
RED = RGBColor(180, 50, 50)
ORANGE = RGBColor(200, 120, 0)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Background
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK_BLUE

    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Inside the Battery"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.2), Inches(8), Inches(1))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Decoding Battery Data via the CAN Bus"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(180, 210, 240)
    p2.alignment = PP_ALIGN.CENTER

    # Subtitle line 2
    txBox3 = slide.shapes.add_textbox(Inches(1), Inches(3.9), Inches(8), Inches(0.6))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "VW ID Buzz (MEB Platform)"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(150, 190, 230)
    p3.alignment = PP_ALIGN.CENTER

    # Info
    txBox4 = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1.2))
    tf4 = txBox4.text_frame
    tf4.word_wrap = True
    p4 = tf4.paragraphs[0]
    p4.text = "Bachelor Thesis - THI / CARISSMA"
    p4.font.size = Pt(14)
    p4.font.color.rgb = RGBColor(180, 200, 220)
    p4.alignment = PP_ALIGN.CENTER
    p5 = tf4.add_paragraph()
    p5.text = "Supervisors: Markus Gregor, Prof. Dr. Hans-Georg Schweiger"
    p5.font.size = Pt(12)
    p5.font.color.rgb = RGBColor(150, 170, 190)
    p5.alignment = PP_ALIGN.CENTER
    p6 = tf4.add_paragraph()
    p6.text = "April 2026"
    p6.font.size = Pt(12)
    p6.font.color.rgb = RGBColor(150, 170, 190)
    p6.alignment = PP_ALIGN.CENTER


def make_slide(prs, title_text, content_func):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    # Title bar background
    shape = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0), Inches(10), Inches(1)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    content_func(slide)
    return slide


def add_bullets(slide, left, top, width, height, items, font_size=16, color=BLACK, bold_first=False):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        if isinstance(item, tuple):
            # (bold_part, normal_part)
            run1 = p.add_run()
            run1.text = item[0]
            run1.font.bold = True
            run1.font.size = Pt(font_size)
            run1.font.color.rgb = color
            run2 = p.add_run()
            run2.text = item[1]
            run2.font.size = Pt(font_size)
            run2.font.color.rgb = color
        else:
            p.text = item
            p.font.size = Pt(font_size)
            p.font.color.rgb = color

        p.space_after = Pt(6)

    return txBox


def add_info_box(slide, left, top, width, height, text, bg_color=LIGHT_BLUE, text_color=DARK_BLUE):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(13)
    p.font.color.rgb = text_color
    return shape


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: TITLE
    # =========================================================================
    add_title_slide(prs)

    # =========================================================================
    # SLIDE 2: AGENDA
    # =========================================================================
    def agenda_content(slide):
        items = [
            "1. Research Question & Motivation",
            "2. What is CAN Bus?",
            "3. The Challenge: Getting Battery Data",
            "4. Data Collection Strategy",
            "5. Tools & Methodology",
            "6. Current Findings",
            "7. Next Steps & Timeline",
        ]
        add_bullets(slide, 0.8, 1.3, 8, 5, items, font_size=20, color=DARK_GRAY)

    make_slide(prs, "Agenda", agenda_content)

    # =========================================================================
    # SLIDE 3: RESEARCH QUESTION
    # =========================================================================
    def research_content(slide):
        # Main question
        add_info_box(slide, 0.8, 1.3, 8.4, 1,
            "Research Question: Can we reverse engineer battery parameters from the CAN bus "
            "of a VW ID Buzz without manufacturer documentation?",
            bg_color=RGBColor(230, 243, 255), text_color=DARK_BLUE)

        add_bullets(slide, 0.8, 2.7, 8, 1.5, [
            ("Focus: ", "HV Battery parameters (SOC, cell voltages, temperatures, current)"),
            ("Vehicle: ", "VW ID Buzz (MEB platform, shared with ID.3/ID.4/ID.5)"),
            ("Method: ", "CAN bus data capture + UDS diagnostic protocol + signal reverse engineering"),
        ], font_size=15)

        # Why it matters
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "Why This Matters"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE

        add_bullets(slide, 0.8, 5, 8, 2, [
            "Battery health monitoring for fleet management and used vehicle assessment",
            "Independent diagnostics without proprietary dealer tools",
            "Academic contribution to EV battery data accessibility research",
            "Foundation for predictive battery maintenance algorithms",
        ], font_size=14, color=DARK_GRAY)

    make_slide(prs, "Research Question & Motivation", research_content)

    # =========================================================================
    # SLIDE 4: WHAT IS CAN BUS
    # =========================================================================
    def can_intro_content(slide):
        add_bullets(slide, 0.8, 1.3, 4.5, 3, [
            "Controller Area Network (Bosch, 1986)",
            "Two-wire shared bus (CAN-H, CAN-L)",
            "Any ECU can broadcast to all others",
            "Priority-based: lower CAN ID = higher priority",
            "Speed: 500 kbps (VW standard)",
            "Up to 8 bytes of data per message",
        ], font_size=15, color=DARK_GRAY)

        # CAN Frame diagram as text box
        add_info_box(slide, 0.8, 4.5, 8.4, 1.2,
            "CAN Frame Structure:\n"
            "[SOF] [CAN ID (11 or 29 bit)] [RTR] [DLC] [DATA: 0-8 bytes] [CRC] [ACK] [EOF]",
            bg_color=RGBColor(240, 245, 250), text_color=DARK_GRAY)

        # Right side - key terms
        txBox = slide.shapes.add_textbox(Inches(5.5), Inches(1.3), Inches(4), Inches(3))
        tf = txBox.text_frame
        tf.word_wrap = True

        terms = [
            ("CAN ID: ", "Identifies message type + priority"),
            ("Payload: ", "The actual data (0-8 bytes)"),
            ("DLC: ", "Data Length Code (how many bytes)"),
            ("11-bit: ", "Standard CAN (0x000-0x7FF)"),
            ("29-bit: ", "Extended CAN (for diagnostics)"),
        ]
        for i, (bold, normal) in enumerate(terms):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r1 = p.add_run()
            r1.text = bold
            r1.font.bold = True
            r1.font.size = Pt(14)
            r1.font.color.rgb = MED_BLUE
            r2 = p.add_run()
            r2.text = normal
            r2.font.size = Pt(14)
            r2.font.color.rgb = DARK_GRAY
            p.space_after = Pt(8)

        add_info_box(slide, 0.8, 5.9, 8.4, 0.8,
            "The VW ID Buzz uses multiple CAN buses interconnected through a central gateway ECU. "
            "The OBD-II port connects through the gateway, which filters most internal traffic.",
            bg_color=RGBColor(255, 245, 230), text_color=RGBColor(150, 90, 0))

    make_slide(prs, "What is CAN Bus?", can_intro_content)

    # =========================================================================
    # SLIDE 5: THE CHALLENGE
    # =========================================================================
    def challenge_content(slide):
        # Left: What we can see
        shape1 = slide.shapes.add_shape(
            1, Inches(0.5), Inches(1.3), Inches(4.2), Inches(2.5)
        )
        shape1.fill.solid()
        shape1.fill.fore_color.rgb = RGBColor(255, 235, 235)
        shape1.line.fill.background()
        tf1 = shape1.text_frame
        tf1.word_wrap = True
        tf1.margin_left = Inches(0.15)
        tf1.margin_top = Inches(0.1)
        p = tf1.paragraphs[0]
        p.text = "Passive CAN Recording (Current)"
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RED

        items1 = [
            "Only 7 CAN IDs captured (0x065-0x06F)",
            "All 11-bit standard frames",
            "Network management / heartbeat messages",
            "NO battery data visible",
            "Gateway blocks internal traffic to OBD-II",
        ]
        for item in items1:
            p = tf1.add_paragraph()
            p.text = f"  - {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(4)

        # Right: What we need
        shape2 = slide.shapes.add_shape(
            1, Inches(5.3), Inches(1.3), Inches(4.2), Inches(2.5)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = RGBColor(230, 255, 230)
        shape2.line.fill.background()
        tf2 = shape2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = Inches(0.15)
        tf2.margin_top = Inches(0.1)
        p = tf2.paragraphs[0]
        p.text = "What We Need (Battery Data)"
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = GREEN

        items2 = [
            "108 individual cell voltages",
            "18 temperature sensor points",
            "Pack voltage, current, SOC",
            "Requires UDS requests to BMS",
            "Uses 29-bit extended CAN IDs",
        ]
        for item in items2:
            p = tf2.add_paragraph()
            p.text = f"  - {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(4)

        # Key insight box
        add_info_box(slide, 0.5, 4.2, 9, 1.2,
            "Key Insight: Battery data is NOT broadcast to the OBD-II port. The BMS only sends "
            "battery parameters when asked via UDS diagnostic protocol (service 0x22). We need "
            "to actively request this data using the correct CAN IDs and protocol.",
            bg_color=RGBColor(255, 248, 220), text_color=RGBColor(140, 100, 0))

        # BMS addressing
        add_info_box(slide, 0.5, 5.7, 9, 1.1,
            "BMS ECU Addressing:\n"
            "Request to BMS:  0x17FC007B (29-bit extended)    |    "
            "Response from BMS:  0x17FE007B (29-bit extended)",
            bg_color=RGBColor(230, 240, 255), text_color=DARK_BLUE)

    make_slide(prs, "The Challenge: Getting Battery Data", challenge_content)

    # =========================================================================
    # SLIDE 6: UDS PROTOCOL
    # =========================================================================
    def uds_content(slide):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "UDS (Unified Diagnostic Services) is the standard protocol for requesting diagnostic data from ECUs."
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY

        # Request example
        shape1 = slide.shapes.add_shape(
            1, Inches(0.5), Inches(2.2), Inches(4.3), Inches(2)
        )
        shape1.fill.solid()
        shape1.fill.fore_color.rgb = RGBColor(240, 245, 255)
        shape1.line.fill.background()
        tf1 = shape1.text_frame
        tf1.word_wrap = True
        tf1.margin_left = Inches(0.15)
        tf1.margin_top = Inches(0.1)
        p = tf1.paragraphs[0]
        p.text = "UDS Request (Read SOC)"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = MED_BLUE

        req_lines = [
            "CAN ID: 0x17FC007B",
            "Data: 03 22 02 8C 55 55 55 55",
            "",
            "  03 = 3 bytes follow",
            "  22 = ReadDataByIdentifier",
            "  02 8C = DID for SOC",
            "  55 = padding",
        ]
        for line in req_lines:
            p = tf1.add_paragraph()
            p.text = line
            p.font.size = Pt(11)
            p.font.name = "Consolas"
            p.font.color.rgb = DARK_GRAY

        # Response example
        shape2 = slide.shapes.add_shape(
            1, Inches(5.2), Inches(2.2), Inches(4.3), Inches(2)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = RGBColor(240, 255, 240)
        shape2.line.fill.background()
        tf2 = shape2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = Inches(0.15)
        tf2.margin_top = Inches(0.1)
        p = tf2.paragraphs[0]
        p.text = "UDS Response (SOC = 80%)"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = GREEN

        resp_lines = [
            "CAN ID: 0x17FE007B",
            "Data: 04 62 02 8C C8 55 55 55",
            "",
            "  04 = 4 bytes follow",
            "  62 = positive response",
            "  02 8C = DID echo",
            "  C8 = 200 -> 200/2.5 = 80%",
        ]
        for line in resp_lines:
            p = tf2.add_paragraph()
            p.text = line
            p.font.size = Pt(11)
            p.font.name = "Consolas"
            p.font.color.rgb = DARK_GRAY

        # Key services
        add_info_box(slide, 0.5, 4.5, 9, 2.3,
            "Key UDS Services:\n"
            "0x10 - DiagnosticSessionControl: Open extended session (required first)\n"
            "0x22 - ReadDataByIdentifier: Read a specific parameter by DID (primary tool)\n"
            "0x3E - TesterPresent: Keep the session alive (send every 2-4 seconds)\n"
            "0x27 - SecurityAccess: Unlock protected parameters (if needed)\n\n"
            "Step 1: Open session (02 10 03)  ->  Step 2: Request DIDs (03 22 XX YY)  ->  Step 3: Keep alive (02 3E 00)",
            bg_color=RGBColor(245, 245, 250), text_color=DARK_GRAY)

    make_slide(prs, "UDS Diagnostic Protocol", uds_content)

    # =========================================================================
    # SLIDE 7: BATTERY PARAMETERS
    # =========================================================================
    def battery_params_content(slide):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8), Inches(0.6))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "145 battery DIDs identified from VW MEB UDS PID database"
        p.font.size = Pt(14)
        p.font.color.rgb = GRAY

        # Categories
        categories = [
            ("State & Energy", RGBColor(200, 230, 255), [
                "SOC (State of Charge) - %",
                "HV Pack Voltage - V",
                "HV Pack Current - A",
                "Car Operation Mode",
            ]),
            ("Cell Voltages", RGBColor(200, 255, 200), [
                "108 individual cell voltages",
                "DIDs: 0x1E40 to 0x1EAB",
                "Formula: (raw/1000)+1 = Volts",
                "Single-frame responses (no ISO-TP)",
            ]),
            ("Temperatures", RGBColor(255, 230, 200), [
                "18 temperature sensor points",
                "Max / Min battery temperature",
                "Main battery temperature",
                "Cooling liquid temps",
            ]),
            ("System", RGBColor(230, 220, 255), [
                "Highest/lowest voltage cell #",
                "Circulation pump duty %",
                "DC-DC converter current/voltage",
                "PTC heater current",
            ]),
        ]

        x_positions = [0.3, 2.7, 5.1, 7.5]
        for idx, (cat_title, bg_color, items) in enumerate(categories):
            x = x_positions[idx]
            shape = slide.shapes.add_shape(
                1, Inches(x), Inches(2), Inches(2.2), Inches(3.5)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = bg_color
            shape.line.fill.background()
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.1)
            tf.margin_right = Inches(0.1)
            tf.margin_top = Inches(0.1)

            p = tf.paragraphs[0]
            p.text = cat_title
            p.font.bold = True
            p.font.size = Pt(13)
            p.font.color.rgb = DARK_BLUE
            p.space_after = Pt(8)

            for item in items:
                p = tf.add_paragraph()
                p.text = f"- {item}"
                p.font.size = Pt(10)
                p.font.color.rgb = DARK_GRAY
                p.space_after = Pt(4)

        add_info_box(slide, 0.3, 5.8, 9.4, 0.8,
            "Priority for thesis: SOC + Pack Voltage + Pack Current + Cell Voltages + Temperatures "
            "(these form the core dataset for battery state analysis)",
            bg_color=RGBColor(255, 248, 220), text_color=RGBColor(140, 100, 0))

    make_slide(prs, "Available Battery Parameters (145 DIDs)", battery_params_content)

    # =========================================================================
    # SLIDE 8: DATA COLLECTION STRATEGY
    # =========================================================================
    def strategy_content(slide):
        strategies = [
            ("Strategy A (Primary)", "CANedge2 Transmit List",
             "Configure CANedge2 to send UDS requests AND log responses. "
             "One device does everything. Need to verify 29-bit extended ID support.",
             GREEN, RGBColor(230, 255, 230)),
            ("Strategy B (Fallback)", "External Diag Tool + CANedge2 Logger",
             "Use OBDEleven/VCDS to send requests. CANedge2 logs passively. "
             "Need Y-splitter OBD cable. More hardware but guaranteed to work.",
             ORANGE, RGBColor(255, 245, 230)),
            ("Strategy C (Last Resort)", "Direct CAN Bus Tap",
             "Bypass gateway, tap directly into powertrain CAN bus. "
             "May find broadcast BMS messages. Requires physical access to wiring.",
             RED, RGBColor(255, 235, 235)),
        ]

        for i, (title, subtitle, desc, title_color, bg_color) in enumerate(strategies):
            y = 1.3 + i * 1.8
            shape = slide.shapes.add_shape(
                1, Inches(0.5), Inches(y), Inches(9), Inches(1.5)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = bg_color
            shape.line.fill.background()
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.2)
            tf.margin_top = Inches(0.1)

            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = f"{title}: "
            r.font.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = title_color
            r = p.add_run()
            r.text = subtitle
            r.font.size = Pt(16)
            r.font.color.rgb = DARK_GRAY

            p2 = tf.add_paragraph()
            p2.text = desc
            p2.font.size = Pt(12)
            p2.font.color.rgb = DARK_GRAY
            p2.space_before = Pt(6)

        add_info_box(slide, 0.5, 6.5, 9, 0.6,
            "Decision: Test Strategy A first (tomorrow). If it fails, pivot to Strategy B.",
            bg_color=RGBColor(230, 243, 255), text_color=DARK_BLUE)

    make_slide(prs, "Data Collection Strategy", strategy_content)

    # =========================================================================
    # SLIDE 9: TOOLS & METHODOLOGY
    # =========================================================================
    def tools_content(slide):
        # Hardware
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(4), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "Hardware"
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_BLUE

        add_bullets(slide, 0.8, 1.8, 4, 2, [
            "CANedge2 (CAN logger + transmitter)",
            "VW ID Buzz (MEB platform)",
            "OBD-II connection (gateway access)",
            "SD card for MF4 log storage",
        ], font_size=13, color=DARK_GRAY)

        # Software
        txBox2 = slide.shapes.add_textbox(Inches(5.5), Inches(1.2), Inches(4), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = "Software"
        p2.font.bold = True
        p2.font.size = Pt(18)
        p2.font.color.rgb = DARK_BLUE

        add_bullets(slide, 5.5, 1.8, 4, 2, [
            "Custom CAN Analyzer (Rust/egui)",
            "Python analysis scripts (asammdf)",
            "MF4 log file processing",
            "Signal reverse engineering tools",
        ], font_size=13, color=DARK_GRAY)

        # Methodology
        txBox3 = slide.shapes.add_textbox(Inches(0.8), Inches(4), Inches(8), Inches(0.5))
        tf3 = txBox3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = "Reverse Engineering Methodology"
        p3.font.bold = True
        p3.font.size = Pt(18)
        p3.font.color.rgb = DARK_BLUE

        steps = [
            ("1. Capture: ", "Record CAN traffic under controlled conditions"),
            ("2. Identify: ", "Find relevant CAN IDs and message patterns"),
            ("3. Decode: ", "Map raw bytes to physical values using known formulas"),
            ("4. Validate: ", "Cross-reference with instrument cluster and physical measurements"),
            ("5. Document: ", "Build signal database (DBC file) for all decoded parameters"),
        ]
        add_bullets(slide, 0.8, 4.6, 8, 2.5, steps, font_size=14, color=DARK_GRAY)

    make_slide(prs, "Tools & Methodology", tools_content)

    # =========================================================================
    # SLIDE 10: CURRENT FINDINGS
    # =========================================================================
    def findings_content(slide):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "Passive CAN recordings analyzed - 7 CAN IDs, 9146 frames, ~333 seconds"
        p.font.size = Pt(13)
        p.font.color.rgb = GRAY

        # Decoded signals table as text
        shape = slide.shapes.add_shape(
            1, Inches(0.5), Inches(1.9), Inches(9), Inches(3.5)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(245, 248, 255)
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = "Decoded Signals from Passive CAN Data"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_BLUE
        p.space_after = Pt(8)

        signals = [
            ("0x065 Byte[0]", "NM_STATE (heartbeat)", "HIGH confidence"),
            ("0x066 Byte[0]", "OperatingMode", "HIGH confidence"),
            ("0x066 Byte[1:2]", "16-bit Counter (big-endian)", "HIGH confidence"),
            ("0x066 Byte[3]", "Temperature_A (raw-40 = 8-10C)", "MEDIUM - maybe battery temp?"),
            ("0x067 Byte[5]", "Temperature_C (raw-40 = 11-13C)", "MEDIUM - maybe battery temp?"),
            ("0x068 Byte[3]", "StatusFlag (binary 0/1)", "HIGH confidence"),
            ("0x06F", "8 independent signals", "LOW - needs more data"),
        ]
        for can_id, signal, conf in signals:
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = f"{can_id}: "
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.name = "Consolas"
            r.font.color.rgb = MED_BLUE
            r = p.add_run()
            r.text = f"{signal}  [{conf}]"
            r.font.size = Pt(11)
            r.font.color.rgb = DARK_GRAY
            p.space_after = Pt(4)

        add_info_box(slide, 0.5, 5.7, 9, 1,
            "Key Finding: Passive OBD-II recording does NOT capture battery data. "
            "The gateway filters internal BMS messages. Active UDS requests are required "
            "to extract battery parameters through the OBD-II port.",
            bg_color=RGBColor(255, 240, 240), text_color=RED)

    make_slide(prs, "Current Findings", findings_content)

    # =========================================================================
    # SLIDE 11: NEXT STEPS
    # =========================================================================
    def next_steps_content(slide):
        phases = [
            ("Week 4 (Now)", "Test UDS communication with BMS via CANedge2. "
             "Verify extended CAN ID support. Extract first battery parameters (SOC, voltage, current).",
             RGBColor(200, 230, 255)),
            ("Weeks 5-6", "Systematic data collection: all 145 battery DIDs. "
             "Record under different conditions: standstill, driving, charging. "
             "Build complete signal database.",
             RGBColor(220, 240, 220)),
            ("Weeks 7-8", "Deep analysis: cell voltage balancing patterns, "
             "temperature distribution, SOC vs voltage curves. "
             "Compare with BMS models from literature.",
             RGBColor(255, 240, 220)),
            ("Weeks 9-12", "Validation, thesis writing, and final presentation. "
             "Create DBC file for community use.",
             RGBColor(240, 235, 255)),
        ]

        for i, (week, desc, bg_color) in enumerate(phases):
            y = 1.3 + i * 1.45
            shape = slide.shapes.add_shape(
                1, Inches(0.5), Inches(y), Inches(9), Inches(1.2)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = bg_color
            shape.line.fill.background()
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.2)
            tf.margin_top = Inches(0.1)

            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = f"{week}: "
            r.font.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = DARK_BLUE
            r = p.add_run()
            r.text = desc
            r.font.size = Pt(13)
            r.font.color.rgb = DARK_GRAY

        add_info_box(slide, 0.5, 6.5, 9, 0.6,
            "Immediate Goal: Tomorrow's test determines if Strategy A works or if we pivot to Strategy B",
            bg_color=RGBColor(255, 248, 220), text_color=RGBColor(140, 100, 0))

    make_slide(prs, "Next Steps & Timeline", next_steps_content)

    # =========================================================================
    # SLIDE 12: THANK YOU
    # =========================================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK_BLUE

    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.8), Inches(8), Inches(1))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Questions?"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(180, 210, 240)
    p2.alignment = PP_ALIGN.CENTER

    txBox3 = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Inside the Battery: Decoding Battery Data via the CAN Bus"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(150, 180, 210)
    p3.alignment = PP_ALIGN.CENTER
    p4 = tf3.add_paragraph()
    p4.text = "THI / CARISSMA - April 2026"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(150, 180, 210)
    p4.alignment = PP_ALIGN.CENTER

    # SAVE
    output_path = OUTPUT_DIR / "Battery_CAN_Research_Presentation.pptx"
    prs.save(str(output_path))
    print(f"Presentation saved to: {output_path}")
    return output_path


if __name__ == "__main__":
    build_presentation()
